#!/usr/bin/env python3
"""Converts the TELx daily report workbook into src/data/report-history.json.

The report's daily sheets hold the program's figures from before the app recorded its own history:
- "Pool Data": one row per pool per day (liquidity, staked share and liquidity, APRs, volume, fees);
- "TELx Data": one row per day for the program, with TEL's price.

Usage:
    python scripts/import-report-history.py "<path to the workbook .xlsx>" [output.json]

Requires openpyxl. The workbook itself is never committed; only this script and its output are.
Cells that are empty, formula errors (#DIV/0!, #N/A, ...) or not finite numbers become null, never 0.
A TEL price of 0 is treated as missing. Rows with no figures at all (dates the report hadn't reached) are
skipped. The summary printed at the end gives the counts.
"""

import json
import math
import sys
from datetime import datetime, timezone

import openpyxl

DEFAULT_OUTPUT = "src/data/report-history.json"

# Workbook pool names to pools. `address` is the pool's entry in src/data/pool.json when exactly one entry
# matches; `candidates` lists the entries when more than one could be meant, and `address` is then null.
POOLS = {
    "TEL/WETH": {
        "key": "balancer-tel-weth",
        "name": "TEL/WETH",
        "label": "TEL 80 WETH 20",
        "chain": "polygon",
        "protocol": "balancer",
        "address": "0xca6efa5704f1ae445e0ee24d9c3ddde34c5be1c2",
    },
    "TEL/BAL": {
        "key": "balancer-tel-bal",
        "name": "TEL/BAL",
        "label": "TEL 80 BAL 20",
        "chain": "polygon",
        "protocol": "balancer",
        "address": "0xa0ef0f4f24662050b7e504d1015a8cd7d2a8a51f",
    },
    "TEL/WPOL": {
        "key": "balancer-tel-wpol",
        "name": "TEL/WPOL",
        "label": "TEL 80 WMATIC 20",
        "chain": "polygon",
        "protocol": "balancer",
        "address": "0x19127998126b6c7d32d596ee15d32719b1789c8e",
    },
    "TEL/USDC": {
        "key": "balancer-tel-usdc",
        "name": "TEL/USDC",
        "label": "TEL 80 USDC 20",
        "chain": "polygon",
        "protocol": "balancer",
        "address": None,
        "candidates": ["0x3bd8a254163f8328efcc4f8c36da566753462433", "0x5c6ee304399dbdb9c8ef030ab642b10820db8f56"],
    },
    "TEL/WBTC": {
        "key": "balancer-tel-wbtc",
        "name": "TEL/WBTC",
        "label": "TEL 80 WBTC 20",
        "chain": "polygon",
        "protocol": "balancer",
        "address": "0xe1e09ce7aac2740846d9b6d9d56f588c65314ecb",
    },
    "TEL/DFX/USDC": {
        "key": "balancer-tel-dfx-usdc",
        "name": "TEL/DFX/USDC",
        "label": "TEL 40 DFX 40 USDC 20",
        "chain": "polygon",
        "protocol": "balancer",
        "address": None,
        "candidates": ["0xc260f3c5a57caf193d1813d8fd0a02442073d6fa", "0x2dbc9ab0160087ae59474fb7bed95b9e808fa6bc"],
    },
    "TEL/WETH (POL)": {
        "key": "uniswap-v4-2025-polygon-tel-weth",
        "name": "TEL/WETH",
        "label": "TEL/WETH, Uniswap v4 (2025)",
        "chain": "polygon",
        "protocol": "uniswap",
        "address": None,
        "candidates": [
            "0x25412ca33f9a2069f0520708da3f70a7843374dd46dc1c7e62f6d5002f5f9fa7",
            "0x9a005a0c12cc2ef01b34e9a7f3fb91a0e6304d377b5479bd3f08f8c29cdf5deb",
        ],
    },
    "TEL/WETH (BASE)": {
        "key": "uniswap-v4-2025-base-tel-eth",
        "name": "TEL/ETH",
        "label": "TEL/ETH, Uniswap v4 (2025)",
        "chain": "base",
        "protocol": "uniswap",
        "address": None,
        "candidates": [
            "0x727b2741ac2b2df8bc9185e1de972661519fc07b156057eeed9b07c50e08829b",
            "0xb6d004fca4f9a34197862176485c45ceab7117c86f07422d1fe3d9cfd6e9d1da",
        ],
    },
    "eMXN/USDC": {
        "key": "uniswap-v4-2025-polygon-usdc-emxn",
        "name": "USDC/eMXN",
        "label": "USDC/eMXN, Uniswap v4 (2025)",
        "chain": "polygon",
        "protocol": "uniswap",
        "address": None,
        "candidates": [
            "0x29f94ec9b66df7fe4068e2d7e9bf0147b49afcdc7cd3283dff03088b8026169f",
            "0xfd56605f7f4620ab44dfc0860d70b9bd1d1f648a5a74558491b39e816a10b99a",
        ],
    },
}

# Column order of each day row in the output, after the day itself (UTC day start, unix seconds).
POOL_FIELDS = ["tvlUSD", "stakedShare", "stakedUSD", "incentivesApr", "volumeUSD", "feesUSD", "feesApr", "totalApr"]
PROGRAM_FIELDS = POOL_FIELDS + ["telUSD"]

# Decimal places per field: whole dollars for levels and volume, cents for fees, five places for fractions.
DECIMALS = {
    "tvlUSD": 0,
    "stakedUSD": 0,
    "volumeUSD": 0,
    "feesUSD": 2,
    "stakedShare": 5,
    "incentivesApr": 5,
    "feesApr": 5,
    "totalApr": 5,
}


class Counts:
    def __init__(self):
        self.dropped = 0
        self.empty_rows = 0
        self.unknown_pools = {}


def number(value, field, counts):
    """The cell as a rounded float, or None (counted as dropped) when it isn't a finite number."""
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        counts.dropped += 1
        return None
    if field == "telUSD":
        if value <= 0:
            counts.dropped += 1
            return None
        return float(f"{value:.6g}")
    rounded = round(float(value), DECIMALS[field])
    return int(rounded) if DECIMALS[field] == 0 else rounded


def day_of(value):
    """UTC day start in unix seconds for a date cell, or None when the cell isn't a date."""
    if not isinstance(value, datetime):
        return None
    return int(datetime(value.year, value.month, value.day, tzinfo=timezone.utc).timestamp())


def has_figures(cells):
    return any(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) for v in cells)


def convert_pool_rows(rows, counts):
    """`rows` are the "Pool Data" rows after the header: Date, Pool, then the POOL_FIELDS columns in order."""
    pools = {}
    for row in rows:
        day = day_of(row[0])
        name = row[1]
        cells = list(row[2 : 2 + len(POOL_FIELDS)])
        if day is None or not isinstance(name, str):
            continue
        if not has_figures(cells):
            counts.empty_rows += 1
            continue
        mapping = POOLS.get(name.strip())
        if mapping is None:
            counts.unknown_pools[name] = counts.unknown_pools.get(name, 0) + 1
            continue
        values = [number(cell, field, counts) for cell, field in zip(cells, POOL_FIELDS)]
        pools.setdefault(mapping["key"], {}).setdefault(day, values)
    return pools


def convert_program_rows(rows, counts):
    """`rows` are the "TELx Data" rows after the header: Date, then the PROGRAM_FIELDS columns in order."""
    days = {}
    for row in rows:
        day = day_of(row[0])
        cells = list(row[1 : 1 + len(PROGRAM_FIELDS)])
        if day is None:
            continue
        # A row with only a 0 TEL price is a date the report hadn't reached.
        if not has_figures(cells[:-1]) and not (isinstance(cells[-1], (int, float)) and cells[-1] > 0):
            counts.empty_rows += 1
            continue
        days.setdefault(day, [number(cell, field, counts) for cell, field in zip(cells, PROGRAM_FIELDS)])
    return days


def build(pool_rows, program_rows):
    counts = Counts()
    pools = convert_pool_rows(pool_rows, counts)
    program = convert_program_rows(program_rows, counts)
    all_days = [day for days in pools.values() for day in days] + list(program)
    out_pools = []
    for mapping in POOLS.values():
        days = pools.get(mapping["key"])
        if not days:
            continue
        entry = {key: mapping[key] for key in ("key", "name", "label", "chain", "protocol", "address")}
        if mapping.get("candidates"):
            entry["candidates"] = mapping["candidates"]
        entry["days"] = [[day, *days[day]] for day in sorted(days)]
        out_pools.append(entry)
    result = {
        "source": "TELx daily report",
        "from": min(all_days) if all_days else None,
        "to": max(all_days) if all_days else None,
        "poolFields": ["day", *POOL_FIELDS],
        "programFields": ["day", *PROGRAM_FIELDS],
        "pools": out_pools,
        "program": [[day, *program[day]] for day in sorted(program)],
    }
    return result, counts


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    output = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT
    workbook = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    pool_rows = list(workbook["Pool Data"].iter_rows(min_row=2, values_only=True))
    program_rows = list(workbook["TELx Data"].iter_rows(min_row=2, values_only=True))
    result, counts = build(pool_rows, program_rows)
    with open(output, "w", encoding="utf-8", newline="\n") as handle:
        json.dump(result, handle, separators=(",", ":"))
        handle.write("\n")
    pool_days = sum(len(pool["days"]) for pool in result["pools"])
    print(f"pools: {len(result['pools'])}, pool days: {pool_days}, program days: {len(result['program'])}")
    print(f"days: {datetime.fromtimestamp(result['from'], timezone.utc).date()} to {datetime.fromtimestamp(result['to'], timezone.utc).date()}")
    print(f"cells dropped as null: {counts.dropped}, rows without figures skipped: {counts.empty_rows}")
    if counts.unknown_pools:
        print(f"rows for unmapped pools skipped: {counts.unknown_pools}")


if __name__ == "__main__":
    main()
